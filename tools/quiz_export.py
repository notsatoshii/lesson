#!/usr/bin/env python3
"""Export a lesson plan's check and retrieval items to Kahoot, Google Forms and Anki.

Usage:
  python quiz_export.py <plan.md> [--kahoot out.xlsx] [--forms out.json] [--anki out.txt]
                        [--title "..."] [--points 1] [--list] [--json]

Item format, written inside a segment's detail block (usually a check or retrieval beat):

  Q: <question>
  A: <correct answer>            one or more A lines; two or more makes a pick-all question
  W: <wrong option> (misconception: <the wrong belief this option catches>)
  W: <wrong option> (오개념: <...>)
  T: 20                          optional, seconds (Kahoot accepts 5, 10, 20, 30, 60, 90, 120, 240)
  F: <feedback shown after answering, says what to do next>     optional

or on one line:  Q: <question> / A: <answer> / W: <wrong> (misconception: ...) / W: <wrong>

Lines may carry a list marker ("- Q: ..."). An item ends at a blank line, the next Q:, or a heading.
Each item is tagged with the segment whose `#### S<n>` heading (or table row) it sits under.

Outputs (only the ones asked for):
  --kahoot  .xlsx in the layout of Kahoot's import template (title rows, header on row 8, data from
            row 9, columns B-H). Enforces questions <= 95 characters and answers <= 60, 2 to 4 answers,
            at least one correct. Items that break a limit are left out and listed; exit 1.
  --forms   a Google Forms API batchUpdate body (isQuiz on, one graded RADIO or CHECKBOX item per
            question, feedback from F: and the misconceptions). Post it to
            forms.batchUpdate on a form you created; <out>.create.json holds the forms.create body.
  --anki    a tab-separated Anki import file with file headers (#separator:tab, #html:true,
            #columns, #deck, #tags column). Front = question, back = answer, wrong options and
            the misconception each one catches, then feedback.

Characters are counted after Unicode NFC normalisation, so one Hangul syllable counts as one.
"""
import argparse
import hashlib
import html
import json
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

Q_MAX, A_MAX = 95, 60
KAHOOT_TIMES = [5, 10, 20, 30, 60, 90, 120, 240]
MISC = re.compile(r"\s*\((?:misconception|오개념|오해)\s*[:：]\s*(.+?)\)\s*$", re.I)
FIELD = re.compile(r"^\s*(?:[-*]\s+|\d+[.)]\s+)?([QAWTF])\s*[:：]\s*(.*)$")


def nfc(s):
    return unicodedata.normalize("NFC", re.sub(r"\s+", " ", s or "").strip())


def strip_md(s):
    s = re.sub(r"\*\*|__|`", "", s)
    m = re.match(r'^["“](.*)["”]$', s.strip())
    return nfc(m.group(1) if m else s)


def parse_wrong(raw):
    raw = strip_md(raw)
    m = MISC.search(raw)
    if m:
        return {"text": nfc(raw[: m.start()]), "misconception": nfc(m.group(1))}
    return {"text": raw, "misconception": ""}


def parse_plan(text):
    title, slug = "", ""
    items, cur, seg = [], None, ""

    def close():
        nonlocal cur
        if cur and cur["question"]:
            items.append(cur)
        cur = None

    for n, line in enumerate(text.splitlines(), 1):
        if not title:
            m = re.match(r"^#\s+(.+)$", line)
            if m:
                title = nfc(m.group(1))
        if not slug:
            m = re.search(r"\bslug\s*:\s*([\w.-]+)", line)
            if m and "<" not in m.group(1):
                slug = m.group(1)
        m = re.match(r"^#{2,6}\s+(S\d+)\b", line)
        if m:
            close()
            seg = m.group(1)
            continue
        if re.match(r"^#{1,6}\s", line):
            close()
            seg = ""
            continue
        m = re.match(r"^\s*\|\s*(S\d+)\s*\|", line)
        if m:
            seg = m.group(1)
        if not line.strip():
            close()
            continue
        # one-line form: Q: ... / A: ... / W: ...
        if re.search(r"\bQ\s*[:：].*\s/\s*A\s*[:：]", line):
            close()
            body = re.sub(r"^\s*(?:[-*]\s+|\d+[.)]\s+)?", "", line)
            cur = {"segment": seg, "line": n, "question": "", "answers": [], "wrong": [], "time": None, "feedback": ""}
            for part in re.split(r"\s+/\s+(?=[QAWTF]\s*[:：])", body):
                fm = re.match(r"^([QAWTF])\s*[:：]\s*(.*)$", part.strip())
                if fm:
                    add_field(cur, fm.group(1), fm.group(2))
            close()
            continue
        m = FIELD.match(line)
        if not m:
            continue
        key, val = m.group(1), m.group(2)
        if key == "Q":
            close()
            cur = {"segment": seg, "line": n, "question": "", "answers": [], "wrong": [], "time": None, "feedback": ""}
        if cur is None:
            continue
        add_field(cur, key, val)
    close()
    return title, slug, items


def add_field(cur, key, val):
    if key == "Q":
        cur["question"] = strip_md(val)
    elif key == "A":
        cur["answers"].append(strip_md(val))
    elif key == "W":
        cur["wrong"].append(parse_wrong(val))
    elif key == "T":
        m = re.search(r"\d+", val)
        cur["time"] = int(m.group(0)) if m else None
    elif key == "F":
        cur["feedback"] = strip_md(val)


def misc_label(text):
    """Label that keeps a misconception from reading like an endorsed fact."""
    return "흔한 오해" if re.search(r"[가-힣]", text) else "Common wrong idea"


def options_for(item):
    """Answer options in a stable shuffled order (so the correct one is not always first)."""
    opts = [(a, True, "") for a in item["answers"]] + [(w["text"], False, w["misconception"]) for w in item["wrong"]]
    key = lambda o: hashlib.sha1((item["question"] + "\x00" + o[0]).encode("utf-8")).hexdigest()
    return sorted(opts, key=key)


def item_id(item, k):
    return f"{item['segment'] or 'Q'}#{k}"


# ---------- Kahoot ----------

def kahoot_problems(item):
    p = []
    if len(item["question"]) > Q_MAX:
        p.append(f"question is {len(item['question'])} characters, Kahoot allows {Q_MAX}")
    n = len(item["answers"]) + len(item["wrong"])
    if not item["answers"]:
        p.append("no A: line (Kahoot needs at least one correct answer)")
    if n < 2:
        p.append("fewer than 2 options")
    if n > 4:
        p.append(f"{n} options, Kahoot allows 4")
    for text, _, _ in options_for(item):
        if len(text) > A_MAX:
            p.append(f"answer '{text[:20]}...' is {len(text)} characters, Kahoot allows {A_MAX}")
    if item["time"] is not None and item["time"] not in KAHOOT_TIMES:
        p.append(f"time {item['time']} s is not one of {KAHOOT_TIMES}")
    return p


def ensure_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("openpyxl not installed, running: pip install openpyxl", file=sys.stderr)
        subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl"], check=True)


def write_kahoot(items, path, title, default_time):
    ensure_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "Quiz template"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B3"] = f"{title} (generated by quiz_export.py)"
    ws["B4"] = "Add questions, at least two answer alternatives, time limit and choose correct answers (at least one)."
    ws["B5"] = "Have fun creating your awesome quiz!"
    ws["B6"] = "Remember: questions have a limit of 95 characters and answers can have 60 characters max."
    header = ["", f"Question - max {Q_MAX} characters", f"Answer 1 - max {A_MAX} characters",
              f"Answer 2 - max {A_MAX} characters", f"Answer 3 - max {A_MAX} characters",
              f"Answer 4 - max {A_MAX} characters",
              "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs",
              "Correct answer(s) - choose at least one"]
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=8, column=c, value=h or None)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for k, item in enumerate(items, 1):
        opts = options_for(item)
        row = 8 + k
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=item["question"])
        for j in range(4):
            ws.cell(row=row, column=3 + j, value=opts[j][0] if j < len(opts) else None)
        ws.cell(row=row, column=7, value=item["time"] or default_time)
        ws.cell(row=row, column=8, value=",".join(str(j + 1) for j, o in enumerate(opts) if o[1]))
    widths = {"A": 4, "B": 48, "C": 26, "D": 26, "E": 26, "F": 26, "G": 16, "H": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wb.save(path)


# ---------- Google Forms ----------

def write_forms(items, path, title, points):
    requests = [{
        "updateSettings": {"settings": {"quizSettings": {"isQuiz": True}}, "updateMask": "quizSettings.isQuiz"}
    }]
    for k, item in enumerate(items):
        opts = options_for(item)
        many = len(item["answers"]) > 1
        wrong_lines = [f"{misc_label(w['misconception'])}: {w['misconception']} ({w['text']})"
                       for w in item["wrong"] if w["misconception"]]
        when_wrong = " ".join(x for x in [item["feedback"], " / ".join(wrong_lines)] if x)
        grading = {"pointValue": points, "correctAnswers": {"answers": [{"value": a} for a in item["answers"]]}}
        if item["feedback"]:
            grading["whenRight"] = {"text": item["feedback"]}
        if when_wrong:
            grading["whenWrong"] = {"text": when_wrong}
        requests.append({"createItem": {"item": {
            "title": item["question"],
            "questionItem": {"question": {
                "required": True,
                "grading": grading,
                "choiceQuestion": {"type": "CHECKBOX" if many else "RADIO",
                                   "options": [{"value": o[0]} for o in opts], "shuffle": True},
            }},
        }, "location": {"index": k}}})
    body = {"requests": requests}
    Path(path).write_text(json.dumps(body, ensure_ascii=False, indent=1), encoding="utf-8")
    create = Path(path).with_suffix(".create.json")
    create.write_text(json.dumps({"info": {"title": title, "documentTitle": title}}, ensure_ascii=False, indent=1),
                      encoding="utf-8")
    return create


# ---------- Anki ----------

def anki_cell(s):
    return html.escape(s, quote=False).replace("\t", " ").replace("\n", "<br>")


def write_anki(items, path, title, slug):
    tag_base = re.sub(r"[^\w-]+", "_", slug or title or "lesson")
    lines = ["#separator:tab", "#html:true", "#notetype:Basic", f"#deck:{title or slug or 'lesson'}",
             "#columns:Front\tBack\tTags", "#tags column:3"]
    for item in items:
        back = [anki_cell(" / ".join(item["answers"]))]
        for w in item["wrong"]:
            line = f"✗ {anki_cell(w['text'])}"
            if w["misconception"]:
                line += f" ({misc_label(w['misconception'])}: {anki_cell(w['misconception'])})"
            back.append(line)
        if item["feedback"]:
            back.append(anki_cell(item["feedback"]))
        tags = f"{tag_base} {tag_base}::{item['segment'] or 'plan'}"
        lines.append(f"{anki_cell(item['question'])}\t{'<br>'.join(back)}\t{tags}")
    Path(path).write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------- main ----------

def main():
    ap = argparse.ArgumentParser(description="Export a plan's Q/A/W items to Kahoot, Google Forms and Anki.")
    ap.add_argument("plan")
    ap.add_argument("--kahoot", help="Kahoot import spreadsheet (.xlsx)")
    ap.add_argument("--forms", help="Google Forms batchUpdate body (.json)")
    ap.add_argument("--anki", help="Anki tab-separated import file (.txt)")
    ap.add_argument("--title", help="quiz title (default: the plan's # heading)")
    ap.add_argument("--time", type=int, default=20, help="default Kahoot time limit in seconds")
    ap.add_argument("--points", type=int, default=1, help="Google Forms points per question")
    ap.add_argument("--list", action="store_true", help="print the parsed items")
    ap.add_argument("--json", action="store_true", help="print the parsed items and checks as JSON")
    a = ap.parse_args()
    if a.time not in KAHOOT_TIMES:
        ap.error(f"--time must be one of {KAHOOT_TIMES}")

    text = Path(a.plan).read_text(encoding="utf-8-sig")
    title, slug, items = parse_plan(text)
    title = a.title or title or slug or "quiz"
    if not items:
        sys.exit(f"{a.plan}: no Q:/A:/W: items found (see the format in this script's docstring)")

    out, exit_code = [], 0
    problems = {item_id(it, k): kahoot_problems(it) for k, it in enumerate(items, 1)}
    for k, it in enumerate(items, 1):
        if not it["answers"]:
            out.append(f"warn: {item_id(it, k)} (line {it['line']}) has no A: line; Forms and Anki get no correct answer")
    if a.list:
        for k, it in enumerate(items, 1):
            out.append(f"{item_id(it, k)} line {it['line']}: Q({len(it['question'])}) {it['question']}")
            for ans in it["answers"]:
                out.append(f"   A({len(ans)}) {ans}")
            for w in it["wrong"]:
                out.append(f"   W({len(w['text'])}) {w['text']}" + (f"  [misconception: {w['misconception']}]" if w["misconception"] else ""))
    if a.kahoot:
        ok = [it for k, it in enumerate(items, 1) if not problems[item_id(it, k)]]
        for key, p in problems.items():
            for x in p:
                out.append(f"kahoot: {key} left out: {x}")
        if len(ok) < len(items):
            exit_code = 1
        if ok:
            write_kahoot(ok, a.kahoot, title, a.time)
            out.append(f"kahoot: {len(ok)} of {len(items)} questions -> {a.kahoot}")
        else:
            out.append("kahoot: no question passes the limits, file not written")
    if a.forms:
        create = write_forms(items, a.forms, title, a.points)
        out.append(f"forms: {len(items)} graded questions -> {a.forms} (forms.batchUpdate body); "
                   f"forms.create body -> {create.name}")
    if a.anki:
        write_anki(items, a.anki, title, slug)
        out.append(f"anki: {len(items)} cards -> {a.anki}")
    if a.json:
        out = [json.dumps({"title": title, "slug": slug, "items": items, "kahoot_problems": problems},
                          ensure_ascii=False, indent=1)]
    if not (a.kahoot or a.forms or a.anki or a.list or a.json):
        out.append(f"{len(items)} items found; pass --kahoot, --forms, --anki or --list")
    sys.stdout.buffer.write(("\n".join(out) + "\n").encode("utf-8"))
    sys.exit(exit_code)


if __name__ == "__main__":
    main()
